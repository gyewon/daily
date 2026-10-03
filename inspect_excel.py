import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\js011\Downloads\2026-09-01~2026-09-27.xlsx', data_only=True)

out = []
for name in wb.sheetnames:
    ws = wb[name]
    sheet_info = {
        "name": name,
        "rows": ws.max_row,
        "cols": ws.max_column,
        "preview": []
    }
    for r in range(1, min(ws.max_row + 1, 15)):
        row_vals = [ws.cell(r, c).value for c in range(1, min(ws.max_column + 1, 15))]
        sheet_info["preview"].append(row_vals)
    out.append(sheet_info)

with open('inspect_excel.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2, default=str)

print("Saved to inspect_excel.json")







