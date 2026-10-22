import os
import sys
import json
import pyzipper
import zipfile
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

with open('banksalad_config.json', 'r', encoding='utf-8') as f:
    cfg = json.load(f)
password = cfg.get('zip_password')

zip_path = 'C:/Users/js011/Downloads/허계원님_2026-10-04~2026-10-22.zip'
out_dir = 'd:/AI App/daily/scratch/extracted_current'
os.makedirs(out_dir, exist_ok=True)
out_excel = os.path.join(out_dir, 'current_upload.xlsx')

extracted = False
pw_bytes = password.encode('utf-8') if password else None

try:
    with pyzipper.AESZipFile(zip_path, 'r') as z:
        if pw_bytes:
            z.setpassword(pw_bytes)
        for fi in z.infolist():
            if fi.filename.endswith('.xlsx') or fi.filename.endswith('.xls'):
                with open(out_excel, 'wb') as f_out:
                    f_out.write(z.read(fi))
                extracted = True
                print(f"Extracted: {fi.filename}")
                break
except Exception as e:
    print(f"AESZipFile error: {e}")

if not extracted:
    try:
        with zipfile.ZipFile(zip_path, 'r') as z:
            if pw_bytes:
                z.setpassword(pw_bytes)
            for fi in z.infolist():
                if fi.filename.endswith('.xlsx') or fi.filename.endswith('.xls'):
                    with open(out_excel, 'wb') as f_out:
                        f_out.write(z.read(fi))
                    extracted = True
                    print(f"Extracted zipfile: {fi.filename}")
                    break
    except Exception as e:
        print(f"Zipfile error: {e}")

if not extracted or not os.path.exists(out_excel):
    print("Failed to extract Excel.")
    sys.exit(1)

wb = openpyxl.load_workbook(out_excel, data_only=True)
ws = wb['가계부 내역'] if '가계부 내역' in wb.sheetnames else wb.active
rows = list(ws.iter_rows(values_only=True))

header = rows[0]
print("Header:", header)

expenses = []
for idx, r in enumerate(rows[1:], start=2):
    if not any(r):
        continue
    # col 2: 구분 ('지출' or '수입' or '이체')
    rec_type = str(r[2] or '').strip()
    if rec_type != '지출':
        continue
    d_raw = str(r[0] or '').split(' ')[0].strip()
    tm_raw = str(r[1] or '').strip()
    cat = str(r[3] or '').strip()
    sub = str(r[4] or '').strip()
    merchant = str(r[5] or '').strip()
    try:
        amt_raw = float(r[6] or 0)
        amt = int(amt_raw)
    except:
        amt = 0
    orig_pay = str(r[8] or '').strip() if len(r) > 8 else ''
    memo = str(r[9] or '').strip() if len(r) > 9 else ''
    
    expenses.append({
        'row': idx,
        'date': d_raw,
        'time': tm_raw,
        'merchant': merchant,
        'amount': amt,
        'category': cat,
        'subCategory': sub,
        'origPay': orig_pay,
        'memo': memo
    })

print(f"Total expense rows in zip Excel: {len(expenses)}")

# Check negative amounts (cancellations/hedging)
negatives = [e for e in expenses if e['amount'] < 0]
positives = [e for e in expenses if e['amount'] > 0]

print(f"\n--- Negative rows ({len(negatives)}) ---")
for neg in negatives:
    print(neg)
    # find matching positive
    matches = [p for p in positives if p['merchant'].strip() == neg['merchant'].strip() and p['amount'] == abs(neg['amount'])]
    print("Matching positive rows:", matches)

# Now inspect which records were added and which were duplicates
# Let's save all expenses to scratch/uploaded_expenses.json
with open('d:/AI App/daily/scratch/uploaded_expenses.json', 'w', encoding='utf-8') as f:
    json.dump(expenses, f, ensure_ascii=False, indent=2)

print("\nSaved expenses to d:/AI App/daily/scratch/uploaded_expenses.json")
