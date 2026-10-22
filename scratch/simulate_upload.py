import os
import sys
import json
import pyzipper
import openpyxl
import datetime

sys.stdout.reconfigure(encoding='utf-8')

with open('banksalad_config.json', 'r', encoding='utf-8') as f:
    cfg = json.load(f)
password = cfg.get('zip_password')

zip_path = 'C:/Users/js011/Downloads/허계원님_2026-10-04~2026-10-22.zip'
out_excel = 'd:/AI App/daily/scratch/extracted_current/current_upload.xlsx'

wb = openpyxl.load_workbook(out_excel, data_only=True)
ws = wb['가계부 내역'] if '가계부 내역' in wb.sheetnames else wb.active
rows = list(ws.iter_rows(values_only=True))

header = [str(c or '').strip() for c in rows[0]]

def find_col(kw_list):
    for idx, name in enumerate(header):
        for kw in kw_list:
            if kw in name:
                return idx
    return -1

col_date = find_col(['날짜', '일자', 'date'])
col_time = find_col(['시간', 'time'])
col_type = find_col(['구분', '유형', 'type'])
col_cat = find_col(['대분류', '카테고리', 'category'])
col_sub = find_col(['소분류', '서브카테고리', 'subcategory'])
col_merchant = find_col(['내용', '가맹점', '사용처', '항목', 'merchant'])
col_amt = find_col(['금액', '지출', 'amount', '결제금액'])
col_pay = find_col(['결제수단', '카드', '결제방법', '자산'])
col_memo = find_col(['메모', '비고', 'memo'])

print(f"Cols -> date:{col_date}, time:{col_time}, type:{col_type}, merchant:{col_merchant}, amt:{col_amt}, pay:{col_pay}")

# Parse rows exactly like app.js
parsed_list = []
for r_idx in range(1, len(rows)):
    row = rows[r_idx]
    if not row or not any(row):
        continue
    if row[0] == '총 합 계' or row[0] == '합계':
        continue
    
    t = str(row[col_type] or '').strip() if col_type != -1 else ''
    if t and '수입' in t:
        continue
        
    raw_amt = row[col_amt]
    if raw_amt is None or raw_amt == '':
        continue
    
    # In JS: const numStr = String(rawAmt).replace(/[^0-9.-]/g, '');
    # const amount = -(Number(numStr) || 0);
    import re
    cleaned = re.sub(r'[^0-9.-]', '', str(raw_amt))
    try:
        amount = -int(float(cleaned))
    except:
        continue
    if amount == 0:
        continue
        
    # Date
    date_raw = row[col_date]
    time_raw = row[col_time] if col_time != -1 else None
    
    if isinstance(date_raw, (datetime.date, datetime.datetime)):
        date_str = date_raw.strftime('%Y-%m-%d')
    else:
        date_str = str(date_raw or '').strip().split(' ')[0].replace('.', '-')
        
    if isinstance(time_raw, datetime.time):
        time_str = time_raw.strftime('%H:%M:%S')
    else:
        time_str = str(time_raw or '00:00:00').strip()
        if len(time_str) == 5:
            time_str += ':00'
            
    merchant = str(row[col_merchant] or '').strip()
    orig_pay = str(row[col_pay] or '기타 카드').strip()
    memo = str(row[col_memo] or '').strip() if col_memo != -1 else ''
    
    # getSignature in JS: (r.date || '') + '|' + (r.time || '') + '|' + (r.merchant || '').trim() + '|' + r.amount + '|' + (r.origPay || r.actualCard || '')
    # Wait, let's check exact getSignature in app.js!
    parsed_list.append({
        'date': date_str,
        'time': time_str,
        'merchant': merchant,
        'amount': amount,
        'origPay': orig_pay,
        'category': str(row[col_cat] or '기타').strip() if col_cat != -1 else '',
        'subCategory': str(row[col_sub] or '').strip() if col_sub != -1 else '',
        'memo': memo,
        'raw_row': r_idx + 1
    })

print(f"Total parsed rows: {len(parsed_list)}")

# Now fetch Supabase current records and deleted signatures
import urllib.request
url_recs = 'https://jkuuwmuniuvijvbtvwde.supabase.co/rest/v1/app_settings?key=eq.records&select=*'
url_del = 'https://jkuuwmuniuvijvbtvwde.supabase.co/rest/v1/app_settings?key=eq.deleted_signatures&select=*'
headers = {
    'apikey': 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImprdXV3bXVuaXV2aWp2YnR2d2RlIiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTA2NjYzNzIsImV4cCI6MjEwNjI0MjM3Mn0.oo9L8jfEKj_ASOqqLCOQaT_8_obkv6OyJ4pGdJXTLuE',
    'Authorization': 'Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImprdXV3bXVuaXV2aWp2YnR2d2RlIiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTA2NjYzNzIsImV4cCI6MjEwNjI0MjM3Mn0.oo9L8jfEKj_ASOqqLCOQaT_8_obkv6OyJ4pGdJXTLuE'
}

req = urllib.request.Request(url_recs, headers=headers)
with urllib.request.urlopen(req) as resp:
    curr_records = json.loads(resp.read().decode('utf-8'))[0]['value']

req2 = urllib.request.Request(url_del, headers=headers)
with urllib.request.urlopen(req2) as resp:
    deleted_sigs = set(json.loads(resp.read().decode('utf-8'))[0]['value'])

print(f"Current DB records: {len(curr_records)}")
print(f"Deleted signatures count: {len(deleted_sigs)}")

# Let's inspect getSignature in JS:
# In app.js:
# function getSignature(r) {
#   return (r.date || '') + '|' + (r.time || '') + '|' + (r.merchant || '').trim() + '|' + r.amount + '|' + (r.origPay || r.actualCard || '');
# }
# Wait, or does getSignature check 4 parts? Let's check getSignature in app.js!
