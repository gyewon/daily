import openpyxl
import os

directory = r'scratch\extracted_excel'
for filename in os.listdir(directory):
    if filename.endswith(".xlsx"):
        filepath = os.path.join(directory, filename)
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb['가계부 내역'] if '가계부 내역' in wb.sheetnames else wb.active
        for r in range(2, ws.max_row+1):
            val = str(ws.cell(r, 6).value)
            if '탕화' in val:
                print(f"Found in {filename}:", ws.cell(r, 1).value, val, ws.cell(r, 7).value)
