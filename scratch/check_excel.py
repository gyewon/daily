import openpyxl

wb = openpyxl.load_workbook(r'scratch\extracted_excel\banksalad_1791221322.xlsx', data_only=True)
ws = wb['가계부 내역'] if '가계부 내역' in wb.sheetnames else wb.active

for r in range(2, ws.max_row+1):
    val = str(ws.cell(r, 6).value)
    if '탕화' in val:
        print("Found:", ws.cell(r, 1).value, val, ws.cell(r, 7).value)
