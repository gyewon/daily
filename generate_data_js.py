import openpyxl
import json
from collections import defaultdict

excel_path = r'C:\Users\js011\Downloads\2026-09-01~2026-09-27.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['가계부 내역']

headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

card_map = {
    '신한은행 The More': '신한은행 The More',
    'KT Plus 우리카드': 'KT Plus 우리카드',
    'MG+ S 하나카드': 'MG+ S 하나카드',
    'KB국민 톡톡 my point카드': 'KB국민 톡톡 my point카드',
    '다드림 LOVE': '신한 복지 다드림 LOVE',
    '신한 복지 다드림 LOVE': '신한 복지 다드림 LOVE',
    '아시아나 KB국민플래티늄카드': '아시아나 KB국민플래티늄카드',
    '네이버페이 간편결제': '[간편결제] 네이버페이',
    '네이버페이 간편결제(포인트)': '[간편결제] 네이버페이(포인트)',
    '카카오페이 간편결제': '[간편결제] 카카오페이',
    '카카오페이 머니': '[간편결제] 카카오페이',
    '페이코 간편결제': '[간편결제] 페이코',
    '토스 간편결제': '[간편결제] 토스',
    '신한 SOL LINK (쏠편한 입출금)': '계좌/현금',
    '7230': '계좌/현금'
}

records = []
card_totals = defaultdict(int)
cat_totals = defaultdict(int)

# 날짜 역순으로 정렬되어 있을 수 있으므로(최신 날짜부터 27일~1일), 날짜 오름차순(1일~27일)으로 정렬할 수도 있음.
# 원본 시트 순서대로 먼저 읽음
raw_records = []
for r in range(2, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if not any(vals):
        continue
    t = str(vals[2] or '').strip()
    if t != '지출':
        continue
    
    d_raw = str(vals[0] or '').split(' ')[0].strip()
    tm_raw = str(vals[1] or '').strip()
    cat = str(vals[3] or '기타').strip()
    sub = str(vals[4] or '').strip()
    merchant = str(vals[5] or '').strip()
    amt = abs(int(vals[6] or 0))
    origPay = str(vals[8] or '기타').strip()
    memo = str(vals[9] or '').strip()
    
    actualCard = card_map.get(origPay, origPay)
    if '입출금' in origPay or '통장' in origPay:
        actualCard = '계좌/현금'
        
    m_str = (str(int(d_raw.split('-')[1])) + '월') if '-' in d_raw and len(d_raw.split('-')) > 1 else '9월'

    raw_records.append({
        'date': d_raw,
        'time': tm_raw,
        'month': m_str,
        'category': cat,
        'subCategory': sub,
        'merchant': merchant,
        'amount': amt,
        'origPay': origPay,
        'actualCard': actualCard,
        'installment': '일시불',
        'billingAmount': amt,
        'exclude': 'N',
        'memo': memo
    })

# 날짜, 시간 순서(오름차순)로 정렬하여 ID 부여
raw_records.sort(key=lambda x: (x['date'], x['time']))

for idx, r in enumerate(raw_records, start=1):
    r['id'] = idx
    r['origId'] = idx
    records.append(r)
    card_totals[r['actualCard']] += r['amount']
    cat_totals[r['category']] += r['amount']

dates = [r['date'] for r in records if r['date']]
period_start = min(dates) if dates else '2026-09-01'
period_end = max(dates) if dates else '2026-09-27'
total_amount = sum(r['amount'] for r in records)

cards_list = [{'name': k, 'amount': v} for k, v in sorted(card_totals.items(), key=lambda x: x[1], reverse=True)]
cats_list = [{'name': k, 'amount': v} for k, v in sorted(cat_totals.items(), key=lambda x: x[1], reverse=True)]

data_obj = {
    "title": "카드별 지출 현황 및 가계부 대시보드",
    "subtitle": f"뱅크샐러드 데이터 기반 ({period_start} ~ {period_end})",
    "period": f"{period_start} ~ {period_end}",
    "physicalCards": [
        "신한은행 The More",
        "KT Plus 우리카드",
        "MG+ S 하나카드",
        "KB국민 톡톡 my point카드",
        "신한 복지 다드림 LOVE",
        "아시아나 KB국민플래티늄카드",
        "기타 카드"
    ],
    "payAndAccounts": [
        "[간편결제] 네이버페이",
        "[간편결제] 네이버페이(포인트)",
        "[간편결제] 카카오페이",
        "[간편결제] 페이코",
        "[간편결제] 토스",
        "계좌/현금"
    ],
    "allCards": [
        "신한은행 The More",
        "KT Plus 우리카드",
        "MG+ S 하나카드",
        "KB국민 톡톡 my point카드",
        "신한 복지 다드림 LOVE",
        "아시아나 KB국민플래티늄카드",
        "기타 카드",
        "[간편결제] 네이버페이",
        "[간편결제] 네이버페이(포인트)",
        "[간편결제] 카카오페이",
        "[간편결제] 페이코",
        "[간편결제] 토스",
        "계좌/현금"
    ],
    "totalRecords": len(records),
    "totalAmount": total_amount,
    "cards": cards_list,
    "categories": cats_list,
    "records": records
}

js_content = "window.INITIAL_DATA = " + json.dumps(data_obj, ensure_ascii=False, indent=2) + ";\n"

with open('data.js', 'w', encoding='utf-8') as f:
    f.write(js_content)

print(f"Generated data.js with {len(records)} records, total {total_amount:,} won.")
