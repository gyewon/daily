"""
Bank Salad to Supabase Auto-Sync Service
뱅크샐러드 이메일(지메일) 첨부 ZIP 자동 다운로드 -> 압축 해제 -> 엑셀 파싱 -> 수파베이스 자동 업로드
"""

import os
import sys
import codecs
if sys.stdout.encoding != 'utf-8':
    sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')
if sys.stderr.encoding != 'utf-8':
    sys.stderr = codecs.getwriter('utf-8')(sys.stderr.buffer, 'strict')
import time
import json
import email
import imaplib
import zipfile
import pyzipper
import openpyxl
import requests
import tkinter as tk
from tkinter import ttk, messagebox
from datetime import datetime
from email.header import decode_header

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, "banksalad_config.json")
PROCESSED_LOG_PATH = os.path.join(BASE_DIR, "processed_mails.json")

SUPABASE_URL = "https://jkuuwmuniuvijvbtvwde.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImprdXV3bXVuaXV2aWp2YnR2d2RlIiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTA2NjYzNzIsImV4cCI6MjEwNjI0MjM3Mn0.oo9L8jfEKj_ASOqqLCOQaT_8_obkv6OyJ4pGdJXTLuE"

SUPABASE_HEADERS = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation"
}

CARD_NAME_MAPPINGS = {
    '신한은행 The More': '신한카드_더모아',
    '신한 The More': '신한카드_더모아',
    '더모아': '신한카드_더모아',
    'KT Plus 우리카드': '우리카드_KT Plus',
    '우리카드 KT Plus': '우리카드_KT Plus',
    'MG+ S 하나카드': '하나카드_MG+ S',
    '하나카드 MG+ S': '하나카드_MG+ S',
    'KB국민 톡톡 my point카드': '국민카드_톡마포',
    '톡톡 my point': '국민카드_톡마포',
    '톡마포': '국민카드_톡마포',
    '다드림 LOVE': '신한카드_다드림 LOVE',
    '신한 복지 다드림 LOVE': '신한카드_다드림 LOVE',
    '신한카드_다드림 LOVE': '신한카드_다드림 LOVE',
    '아시아나 KB국민플래티늄카드': '국민카드_KB국민플래티늄카드',
    'KB국민플래티늄카드': '국민카드_KB국민플래티늄카드',
    '국민카드_KB국민플래티늄카드': '국민카드_KB국민플래티늄카드',
    '현대카드_제로에디션': '현대카드_제로에디션',
    '현대카드 ZERO': '현대카드_제로에디션',
    '네이버페이 간편결제': '네이버페이 간편결제',
    '네이버페이 간편결제(포인트)': '네이버페이 간편결제(포인트)',
    '카카오페이 간편결제': '카카오페이 간편결제',
    '카카오페이 머니': '카카오페이 머니',
    '페이코 간편결제': '페이코 간편결제',
    '토스 간편결제': '토스 간편결제',
    '신한 SOL LINK (쏠편한 입출금)': '신한 SOL LINK (쏠편한 입출금)',
    'Sh평생주거래우대통장(스페셜플러스예금-잔액구간별)': 'Sh평생주거래우대통장(스페셜플러스예금-잔액구간별)',
    '상상모바일통장': '상상모바일통장',
    '7230': '7230',
    '현금': '현금'
}

def load_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            pass
    return {
        "gmail_user": "",
        "gmail_app_password": "",
        "zip_password": "",
        "check_interval_seconds": 60,
        "watch_downloads_folder": False
    }

def save_config(cfg):
    with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)

def load_processed_mails():
    if os.path.exists(PROCESSED_LOG_PATH):
        try:
            with open(PROCESSED_LOG_PATH, 'r', encoding='utf-8') as f:
                return set(json.load(f))
        except Exception:
            pass
    return set()

def save_processed_mails(ids_set):
    with open(PROCESSED_LOG_PATH, 'w', encoding='utf-8') as f:
        json.dump(list(ids_set), f, ensure_ascii=False, indent=2)

def decode_mime_header(header_str):
    if not header_str:
        return ""
    decoded_fragments = decode_header(header_str)
    result = []
    for fragment, encoding in decoded_fragments:
        if isinstance(fragment, bytes):
            if encoding:
                try:
                    result.append(fragment.decode(encoding, errors='replace'))
                except Exception:
                    result.append(fragment.decode('utf-8', errors='replace'))
            else:
                result.append(fragment.decode('utf-8', errors='replace'))
        else:
            result.append(str(fragment))
    return "".join(result)

def extract_excel_from_zip(zip_data_or_path, password):
    """
    Extracts the excel file from Bank Salad password-protected ZIP archive.
    Returns path to extracted temporary .xlsx file.
    """
    tmp_out_dir = os.path.join(BASE_DIR, "scratch", "extracted_excel")
    os.makedirs(tmp_out_dir, exist_ok=True)
    out_file = os.path.join(tmp_out_dir, f"banksalad_{int(time.time())}.xlsx")

    pw_bytes = str(password).strip().encode('utf-8') if password else None
    
    extracted = False
    try:
        with pyzipper.AESZipFile(zip_data_or_path, 'r') as z:
            if pw_bytes:
                z.setpassword(pw_bytes)
            for file_info in z.infolist():
                if file_info.filename.endswith('.xlsx') or file_info.filename.endswith('.xls'):
                    data = z.read(file_info)
                    with open(out_file, 'wb') as f_out:
                        f_out.write(data)
                    extracted = True
                    break
    except Exception as e:
        pass

    if not extracted:
        try:
            with zipfile.ZipFile(zip_data_or_path, 'r') as z:
                if pw_bytes:
                    z.setpassword(pw_bytes)
                for file_info in z.infolist():
                    if file_info.filename.endswith('.xlsx') or file_info.filename.endswith('.xls'):
                        data = z.read(file_info)
                        with open(out_file, 'wb') as f_out:
                            f_out.write(data)
                        extracted = True
                        break
        except Exception as e:
            pass

    if extracted and os.path.exists(out_file):
        return out_file
    return None

def parse_banksalad_excel(xlsx_path):
    """
    Parses Bank Salad Excel file and returns list of raw expense records.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if '가계부 내역' in wb.sheetnames:
        ws = wb['가계부 내역']
    else:
        ws = wb.active

    raw_records = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not any(vals):
            continue
        
        # 구분 (vals[2]) -> '지출'만 처리
        record_type = str(vals[2] or '').strip()
        if record_type != '지출':
            continue
        
        d_raw = str(vals[0] or '').split(' ')[0].strip()
        tm_raw = str(vals[1] or '').strip()
        cat = str(vals[3] or '기타').strip()
        sub = str(vals[4] or '').strip()
        merchant = str(vals[5] or '').strip()
        
        try:
            amt = abs(int(float(vals[6] or 0)))
        except (ValueError, TypeError):
            amt = 0
            
        orig_pay = str(vals[8] or '기타').strip() if len(vals) > 8 else '기타'
        memo = str(vals[9] or '').strip() if len(vals) > 9 else ''
        
        if not d_raw or amt == 0:
            continue
            
        actual_card = CARD_NAME_MAPPINGS.get(orig_pay, orig_pay)
        if '입출금' in orig_pay or '통장' in orig_pay:
            actual_card = '계좌/현금'
            
        m_str = f"{int(d_raw.split('-')[1])}월" if '-' in d_raw and len(d_raw.split('-')) > 1 else '9월'
        
        sig = f"{d_raw}|{tm_raw}|{merchant}|{amt}"
        fuzzy_sig = f"{d_raw}|{merchant}|{amt}"
        
        raw_records.append({
            'date': d_raw,
            'time': tm_raw,
            'month': m_str,
            'category': cat,
            'subCategory': sub,
            'merchant': merchant,
            'amount': amt,
            'origPay': orig_pay,
            'actualCard': actual_card,
            'isInstallment': 'N',
            'installment': '일시불',
            'billingAmount': amt,
            'exclude': 'N',
            'memo': memo,
            'originalSignature': sig,
            '_fuzzy_sig': fuzzy_sig
        })

    return raw_records

def get_supabase_app_settings():
    """
    Fetches all app_settings from Supabase.
    """
    try:
        resp = requests.get(f"{SUPABASE_URL}/rest/v1/app_settings?select=*", headers=SUPABASE_HEADERS, timeout=10)
        if resp.status_code == 200:
            data = resp.json()
            settings = {}
            for item in data:
                settings[item.get('key')] = item.get('value')
            return settings
    except Exception as e:
        print(f"[Supabase] Error loading app_settings: {e}")
    return {}

def apply_rules_to_record(rec, rules):
    """
    Auto-categorize record according to user's registered rules.
    """
    if not rules:
        return
    merchant = rec.get('merchant', '')
    amount = rec.get('amount', 0)
    for rule in rules:
        kw = rule.get('keyword', '').strip()
        rule_amt = rule.get('amount')
        
        kw_match = bool(kw and kw in merchant)
        amt_match = (rule_amt is None or rule_amt == amount or rule_amt == '')
        
        if kw_match and amt_match:
            if rule.get('category'):
                rec['category'] = rule['category']
            if rule.get('subCategory') is not None:
                rec['subCategory'] = rule['subCategory']
            break

def sync_records_to_supabase(new_raw_records):
    """
    Merges newly extracted raw records with existing Supabase records and saves back.
    User's manually edited cards and fields are strictly preserved!
    """
    settings = get_supabase_app_settings()
    existing_records = settings.get('records') or []
    deleted_signatures = set(settings.get('deleted_signatures') or [])
    deleted_fuzzy_sigs = set()
    for ds in deleted_signatures:
        parts = ds.split('|')
        if len(parts) >= 4:
            deleted_fuzzy_sigs.add(f"{parts[0]}|{parts[2]}|{parts[3]}")

    rules = settings.get('rules') or []
    notification_logs = settings.get('notification_logs') or []

    # Map exact signatures & fuzzy signatures of existing records
    existing_sigs = set()
    existing_fuzzy_sigs = set()
    max_id = 0
    for r in existing_records:
        sig = r.get('originalSignature') or f"{r.get('date')}|{r.get('time', '')}|{r.get('merchant')}|{r.get('amount')}"
        existing_sigs.add(sig)
        fuzzy = f"{r.get('date')}|{r.get('merchant')}|{r.get('amount')}"
        existing_fuzzy_sigs.add(fuzzy)
        if r.get('id') and isinstance(r.get('id'), int):
            max_id = max(max_id, r.get('id'))

    added_records = []
    for r in new_raw_records:
        sig = r['originalSignature']
        fuzzy = r['_fuzzy_sig']
        
        # Skip if already deleted by user or already present in Supabase
        if sig in deleted_signatures or fuzzy in deleted_fuzzy_sigs:
            continue
        if sig in existing_sigs or fuzzy in existing_fuzzy_sigs:
            continue
        
        # Clean temporary helper key
        r.pop('_fuzzy_sig', None)
        
        # Apply user's custom category rules
        apply_rules_to_record(r, rules)
        
        max_id += 1
        r['id'] = max_id
        r['origId'] = max_id
        
        added_records.append(r)
        existing_sigs.add(sig)
        existing_fuzzy_sigs.add(fuzzy)

    if not added_records:
        print("[Sync] 신규 추가할 새로운 내역이 없습니다 (이미 등록되어 있거나 중복).")
        return 0

    print(f"[Sync] 신규 지출 {len(added_records)}건을 Supabase에 안전하게 추가 중...")
    all_records = existing_records + added_records

    # Add log entry
    log_msg = f"뱅크샐러드 자동 동기화: 신규 지출 {len(added_records)}건이 추가되었습니다."
    log_entry = {
        "id": int(time.time() * 1000),
        "time": datetime.now().strftime("%Y. %m. %d. %H:%M:%S"),
        "message": log_msg,
        "type": "success"
    }
    notification_logs.insert(0, log_entry)
    if len(notification_logs) > 50:
        notification_logs = notification_logs[:50]

    # Save to app_settings
    payload = [
        {"key": "records", "value": all_records},
        {"key": "notification_logs", "value": notification_logs}
    ]
    
    try:
        resp = requests.post(
            f"{SUPABASE_URL}/rest/v1/app_settings",
            headers={**SUPABASE_HEADERS, "Prefer": "resolution=merge-duplicates"},
            json=payload,
            timeout=15
        )
        if resp.status_code in (200, 201, 204):
            print(f"[Sync] ✅ Supabase 동기화 성공! ({len(added_records)}건 추가)")
            return len(added_records)
        else:
            print(f"[Sync] ❌ Supabase 저장 실패: {resp.status_code} {resp.text}")
    except Exception as e:
        print(f"[Sync] 네트워크 오류: {e}")

    return 0

def fetch_gmail_banksalad_attachments(cfg, check_all=False):
    """
    Connects to Gmail via IMAP, searches for the latest Bank Salad emails,
    downloads and processes attached ZIP files.
    """
    user = cfg.get("gmail_user", "").strip()
    pwd = cfg.get("gmail_app_password", "").strip().replace(" ", "")
    zip_pwd = cfg.get("zip_password", "").strip()

    if not user or not pwd:
        return 0

    processed_ids = load_processed_mails()
    new_records_count = 0

    try:
        mail = imaplib.IMAP4_SSL("imap.gmail.com", 993)
        mail.login(user, pwd)
        mail.select("INBOX")

        search_queries = [
            '(OR SUBJECT "뱅크샐러드" SUBJECT "가계부")',
            '(OR FROM "banksalad" SUBJECT "banksalad")'
        ]

        found_msg_ids = set()
        for q in search_queries:
            try:
                status, data = mail.search(None, q)
                if status == 'OK' and data[0]:
                    found_msg_ids.update(data[0].split())
            except Exception:
                pass

        if not found_msg_ids:
            try:
                status, data = mail.search(None, 'ALL')
                if status == 'OK' and data[0]:
                    all_ids = data[0].split()
                    found_msg_ids.update(all_ids[-5:])
            except Exception:
                pass

        sorted_ids = sorted(list(found_msg_ids), key=lambda x: int(x), reverse=True)
        
        # Only check the newest unprocessed emails (max 2 newest if not check_all)
        target_ids = sorted_ids if check_all else sorted_ids[:2]

        for msg_id in target_ids:
            id_str = msg_id.decode() if isinstance(msg_id, bytes) else str(msg_id)
            if id_str in processed_ids:
                continue

            status, msg_data = mail.fetch(msg_id, '(RFC822)')
            if status != 'OK' or not msg_data:
                continue

            raw_email = msg_data[0][1]
            msg = email.message_from_bytes(raw_email)

            subject = decode_mime_header(msg.get('Subject', ''))
            from_addr = decode_mime_header(msg.get('From', ''))

            # Check attachments
            zip_attachments = []
            for part in msg.walk():
                filename = part.get_filename()
                if filename:
                    decoded_fn = decode_mime_header(filename)
                    if decoded_fn.lower().endswith('.zip'):
                        zip_attachments.append((decoded_fn, part.get_payload(decode=True)))

            zip_success = False
            for fn, zip_bytes in zip_attachments:
                print(f"[Gmail] 최신 메일 첨부파일 처리 중: {fn} (메일 제목: {subject})...")
                tmp_zip = os.path.join(BASE_DIR, "scratch", f"mail_{int(time.time())}.zip")
                os.makedirs(os.path.dirname(tmp_zip), exist_ok=True)
                with open(tmp_zip, 'wb') as f:
                    f.write(zip_bytes)

                xlsx_path = extract_excel_from_zip(tmp_zip, zip_pwd)
                if xlsx_path and os.path.exists(xlsx_path):
                    raw_recs = parse_banksalad_excel(xlsx_path)
                    if raw_recs:
                        count = sync_records_to_supabase(raw_recs)
                        new_records_count += count
                        zip_success = True
                        print(f"[Gmail] ✅ {fn} 파일 처리 및 동기화 성공! (신규 {count}건 반영)")
                else:
                    print(f"[Gmail] ❌ {fn} 압축 해제 실패: 압축 비밀번호(생년월일 6자리)가 맞는지 확인해주세요. (현재 입력값: {zip_pwd})")

            if zip_success:
                processed_ids.add(id_str)
                save_processed_mails(processed_ids)

        mail.logout()
    except Exception as e:
        print(f"[Gmail IMAP Error] {e}")

    return new_records_count

def mark_all_current_mails_as_processed():
    """
    Marks all existing emails in inbox as already processed so old historical
    emails are never crawled again.
    """
    cfg = load_config()
    user = cfg.get("gmail_user", "").strip()
    pwd = cfg.get("gmail_app_password", "").strip().replace(" ", "")
    if not user or not pwd:
        return
    try:
        mail = imaplib.IMAP4_SSL("imap.gmail.com", 993)
        mail.login(user, pwd)
        mail.select("INBOX")
        status, data = mail.search(None, 'ALL')
        if status == 'OK' and data[0]:
            all_ids = set(x.decode() if isinstance(x, bytes) else str(x) for x in data[0].split())
            save_processed_mails(all_ids)
            print(f"[Gmail Init] 기존 메일 {len(all_ids)}건을 처리 완료 목록에 등록했습니다 (과거 메일 중복 수집 방지).")
        mail.logout()
    except Exception as e:
        print(f"[Gmail Init Error] {e}")

def run_single_sync(cfg):
    print(f"\n[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] 뱅크샐러드 최신 메일 확인 시작...")
    cnt = fetch_gmail_banksalad_attachments(cfg, check_all=False)
    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] 동기화 완료: 신규 지출 {cnt}건 반영\n")
    return cnt

def run_daemon_loop():
    cfg = load_config()
    if not cfg.get("gmail_user") or not cfg.get("gmail_app_password"):
        print("[설정 안내] Gmail 계정 또는 앱 비밀번호가 설정되지 않았습니다. 설정 창을 실행합니다...")
        open_gui_settings()
        cfg = load_config()

    # Mark existing emails so daemon only reacts to NEW incoming emails
    mark_all_current_mails_as_processed()

    interval = max(10, cfg.get("check_interval_seconds", 60))
    print(f"\n[실시간 자동 감지 데몬 시작] 감시 주기: {interval}초")
    print("뱅크샐러드에서 새로운 내보내기 메일이 수신되면 자동으로 감지하여 동기화합니다.")
    print("종료하려면 Ctrl + C 를 누르세요.\n")

    while True:
        try:
            cfg = load_config()
            fetch_gmail_banksalad_attachments(cfg, check_all=False)
        except Exception as e:
            print(f"[동기화 오류] {e}")
        time.sleep(interval)

# --- GUI Settings Window ---
def open_gui_settings():
    root = tk.Tk()
    root.title("뱅크샐러드 자동 동기화 설정")
    root.geometry("480x420")
    root.resizable(False, False)
    root.configure(bg="#f8fafc")

    cfg = load_config()

    # Title
    title_frame = tk.Frame(root, bg="#1e293b", padx=16, pady=16)
    title_frame.pack(fill="x")
    
    tk.Label(title_frame, text="⚡ 뱅크샐러드 메일 자동 수집 설정", font=("Pretendard", 13, "bold"), fg="#ffffff", bg="#1e293b").pack(anchor="w")
    tk.Label(title_frame, text="지메일로 수신된 최신 뱅크샐러드 엑셀을 실시간 감지하여 자동 등록합니다.", font=("Pretendard", 9), fg="#94a3b8", bg="#1e293b").pack(anchor="w", pady=(4, 0))

    content = tk.Frame(root, bg="#f8fafc", padx=20, pady=16)
    content.pack(fill="both", expand=True)

    # 1. Gmail Email
    tk.Label(content, text="1. 수신 구글 지메일 주소", font=("Pretendard", 9, "bold"), bg="#f8fafc", fg="#334155").pack(anchor="w", pady=(4, 2))
    ent_user = ttk.Entry(content, width=45)
    ent_user.pack(fill="x", pady=(0, 8))
    ent_user.insert(0, cfg.get("gmail_user", ""))

    # 2. Gmail App Password
    pw_lbl_frame = tk.Frame(content, bg="#f8fafc")
    pw_lbl_frame.pack(fill="x", pady=(4, 2))
    tk.Label(pw_lbl_frame, text="2. 구글 앱 비밀번호 (16자리)", font=("Pretendard", 9, "bold"), bg="#f8fafc", fg="#334155").pack(side="left")
    
    ent_app_pw = ttk.Entry(content, width=45, show="*")
    ent_app_pw.pack(fill="x", pady=(0, 8))
    ent_app_pw.insert(0, cfg.get("gmail_app_password", ""))

    # 3. ZIP Password (Birthdate)
    tk.Label(content, text="3. 뱅크샐러드 압축 해제 비밀번호 (생년월일 6자리)", font=("Pretendard", 9, "bold"), bg="#f8fafc", fg="#334155").pack(anchor="w", pady=(4, 2))
    ent_zip_pw = ttk.Entry(content, width=45, show="*")
    ent_zip_pw.pack(fill="x", pady=(0, 8))
    ent_zip_pw.insert(0, cfg.get("zip_password", ""))

    # 4. Check Interval
    tk.Label(content, text="4. 메일 확인 주기 (초 단위, 기본 60초)", font=("Pretendard", 9, "bold"), bg="#f8fafc", fg="#334155").pack(anchor="w", pady=(4, 2))
    ent_interval = ttk.Entry(content, width=45)
    ent_interval.pack(fill="x", pady=(0, 12))
    ent_interval.insert(0, str(cfg.get("check_interval_seconds", 60)))

    # Status Label
    status_var = tk.StringVar(value="")
    lbl_status = tk.Label(content, textvariable=status_var, font=("Pretendard", 9), fg="#2563eb", bg="#f8fafc")
    lbl_status.pack(anchor="w", pady=(0, 10))

    # Buttons
    btn_frame = tk.Frame(root, bg="#f8fafc", padx=20, pady=12)
    btn_frame.pack(fill="x", side="bottom")

    def test_and_save():
        user = ent_user.get().strip()
        app_pw = ent_app_pw.get().strip().replace(" ", "")
        zip_pw = ent_zip_pw.get().strip()
        try:
            interval = int(ent_interval.get().strip() or "60")
        except ValueError:
            interval = 60

        if not user or not app_pw:
            messagebox.showerror("입력 오류", "지메일 주소와 앱 비밀번호를 입력해주세요.")
            return

        status_var.set("⏳ 지메일 연결 테스트 중...")
        root.update()

        try:
            mail = imaplib.IMAP4_SSL("imap.gmail.com", 993)
            mail.login(user, app_pw)
            mail.logout()
            status_var.set("✅ 지메일 연결 성공! 설정을 저장했습니다.")
            
            cfg["gmail_user"] = user
            cfg["gmail_app_password"] = app_pw
            cfg["zip_password"] = zip_pw
            cfg["check_interval_seconds"] = interval
            save_config(cfg)
            
            # Register current inbox emails so they won't be re-crawled
            mark_all_current_mails_as_processed()
            
            messagebox.showinfo("성공", "지메일 연결 테스트 성공!\n설정이 정상적으로 저장되었습니다.")
            root.destroy()
        except Exception as e:
            status_var.set(f"❌ 연결 실패: {e}")
            messagebox.showerror("연결 실패", f"지메일 로그인에 실패했습니다.\n\n오류: {e}\n\n* 구글 계정 보안 설정에서 2단계 인증 활성화 후 [앱 비밀번호(16자리)]를 발급받아 입력해주세요.")

    btn_save = tk.Button(btn_frame, text="연결 테스트 & 저장", font=("Pretendard", 10, "bold"), bg="#2563eb", fg="#ffffff", padx=14, pady=6, relief="flat", command=test_and_save)
    btn_save.pack(side="right", padx=(6, 0))

    root.mainloop()

if __name__ == "__main__":
    if len(sys.argv) > 1:
        arg = sys.argv[1].lower()
        if arg in ("--gui", "-g", "config", "setup"):
            open_gui_settings()
        elif arg in ("--sync", "-s", "once"):
            run_single_sync(load_config())
        elif arg in ("--init", "-i"):
            mark_all_current_mails_as_processed()
        else:
            run_daemon_loop()
    else:
        cfg = load_config()
        if not cfg.get("gmail_user") or not cfg.get("gmail_app_password"):
            open_gui_settings()
        else:
            run_daemon_loop()
